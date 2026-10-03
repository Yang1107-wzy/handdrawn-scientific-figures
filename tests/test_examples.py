"""Rebuild real public examples in an isolated directory; no network or image tool."""

import csv
import hashlib
import json
from pathlib import Path
import runpy
import shutil
import subprocess
import sys
import tempfile
import unittest

from PIL import Image
from pypdf import PdfReader
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]


class ExampleBuildTests(unittest.TestCase):
    def test_method_feedback_does_not_cross_authored_labels(self):
        api = runpy.run_path(str(ROOT / "examples/build_examples.py"))
        config = api["method_manifest"]()
        for layer in config["layers"]:
            if layer["type"] == "image":
                layer["path"] = str((ROOT / "examples/method" / layer["path"]).resolve())
        with tempfile.TemporaryDirectory() as temp:
            manifest = Path(temp) / "figure.json"
            manifest.write_text(json.dumps(config))
            report = api["compose_figure"](manifest, Path(temp) / "out")
            feedback = [item for item in config["layers"] if item["type"] == "arrow"][-1]
            hits = []
            for a, b in zip(feedback["points"], feedback["points"][1:]):
                for item in report["texts"]:
                    x, y, w, h = item["bounds_mm"]
                    if (
                        a[0] == b[0]
                        and x < a[0] < x + w
                        and min(a[1], b[1]) < y + h
                        and max(a[1], b[1]) > y
                    ):
                        hits.append(item["text"])
                    if (
                        a[1] == b[1]
                        and y < a[1] < y + h
                        and min(a[0], b[0]) < x + w
                        and max(a[0], b[0]) > x
                    ):
                        hits.append(item["text"])
            self.assertFalse(hits, "Feedback intersects authored text: " + repr(hits))

    def test_skill_only_archive_has_its_own_runtime_dependencies(self):
        from zipfile import ZipFile

        with tempfile.TemporaryDirectory() as temp:
            output = Path(temp) / "skill.zip"
            run = subprocess.run(
                [
                    sys.executable,
                    str(ROOT / "scripts/package.py"),
                    "--skill-only",
                    "--output",
                    str(output),
                ],
                capture_output=True,
                text=True,
            )
            self.assertEqual(run.returncode, 0, run.stderr)
            with ZipFile(output) as archive:
                required = "handdrawn-scientific-figures/requirements.txt"
                self.assertIn(required, archive.namelist())
                deps = archive.read(required).decode()
                self.assertIn("matplotlib", deps)
                self.assertIn("python-pptx", deps)

    def test_offline_examples_preserve_art_and_deliver_paper_and_slide(self):
        with tempfile.TemporaryDirectory() as temp:
            copy = Path(temp)
            shutil.copytree(ROOT / "skills", copy / "skills")
            shutil.copytree(ROOT / "examples", copy / "examples")
            originals = {
                p.name: hashlib.sha256(p.read_bytes()).hexdigest()
                for p in (copy / "examples/art").glob("*.png")
            }
            run = subprocess.run(
                [sys.executable, str(copy / "examples/build_examples.py")],
                cwd=copy,
                capture_output=True,
                text=True,
            )
            self.assertEqual(run.returncode, 0, run.stdout + run.stderr)
            for name in ("method", "scene", "hybrid"):
                folder = copy / "examples" / name
                with Image.open(folder / "figure.png") as im:
                    self.assertEqual(im.size, (2126, 1299))
                with Image.open(folder / "slide.png") as im:
                    self.assertEqual(im.size, (3840, 2160))
                pdf = PdfReader(folder / "figure.pdf")
                self.assertEqual(len(pdf.pages), 1)
                self.assertGreater(len(pdf.pages[0].extract_text()), 40)
                self.assertFalse(pdf.is_encrypted)
                ppt = Presentation(folder / "slide.pptx")
                self.assertEqual(len(ppt.slides), 1)
                self.assertTrue(any(s.has_text_frame for s in ppt.slides[0].shapes))
                for suffix in ("figure.svg", "slide.svg"):
                    svg = (folder / suffix).read_text()
                    self.assertIn("<text ", svg)
                    self.assertIn("data:image/png;base64,", svg)
                report = json.loads((folder / "figure.render-report.json").read_text())
                self.assertFalse(report["bounds_warnings"])
                self.assertNotIn("/Users/", (folder / "figure.json").read_text())
            self.assertEqual(
                originals,
                {
                    p.name: hashlib.sha256(p.read_bytes()).hexdigest()
                    for p in (copy / "examples/art").glob("*.png")
                },
            )
            with (copy / "examples/hybrid/teaching_signal.csv").open(newline="") as f:
                rows = list(csv.DictReader(f))
            self.assertEqual(len(rows), 25)
            expected = {0: 1.2, 6: 1.5, 12: 1.2, 18: 0.9, 24: 1.2}
            for index, value in expected.items():
                self.assertAlmostEqual(float(rows[index]["analytic_m"]), value, places=7)
            self.assertIn(
                "Synthetic", PdfReader(copy / "examples/hybrid/figure.pdf").pages[0].extract_text()
            )


if __name__ == "__main__":
    unittest.main()
