"""Offline renderer contract tests; run with python -m unittest discover -s tests."""

import hashlib
import json
from pathlib import Path
import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET

from PIL import Image
from pypdf import PdfReader
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "skills/handdrawn-scientific-figures/scripts"
sys.path.insert(0, str(SCRIPTS))


class RendererTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.base = Path(self.tmp.name)
        self.asset = self.base / "scene.png"
        Image.new("RGB", (600, 300), "#d8e6d1").save(self.asset)
        self.original_hash = hashlib.sha256(self.asset.read_bytes()).hexdigest()
        self.manifest = self.base / "figure.json"
        self.config = {
            "profile": "paper",
            "layers": [
                {"type": "image", "path": "scene.png", "x": 10, "y": 20, "width": 60, "height": 40},
                {
                    "type": "text",
                    "text": "Scientific title",
                    "font": "title",
                    "font_size": 18,
                    "x": 10,
                    "y": 4,
                },
                {
                    "type": "text",
                    "text": "Selectable evidence",
                    "font": "body",
                    "font_size": 11,
                    "x": 90,
                    "y": 20,
                },
                {
                    "type": "math",
                    "text": r"$x_i^2 + \\alpha$".replace("\\\\", "\\"),
                    "x": 90,
                    "y": 40,
                    "font_size": 16,
                },
                {
                    "type": "arrow",
                    "points": [[75, 45], [83, 45], [87, 30]],
                    "color": "#126b70",
                    "width": 1.3,
                },
                {
                    "type": "rect",
                    "x": 6,
                    "y": 17,
                    "width": 67,
                    "height": 47,
                    "stroke": "#126b70",
                    "rough": True,
                },
            ],
        }

    def render(self, config=None, overwrite=False):
        self.assertTrue(
            (SCRIPTS / "compose_figure.py").exists(), "renderer implementation is missing"
        )
        from compose_figure import compose_figure

        self.manifest.write_text(json.dumps(config or self.config))
        return compose_figure(self.manifest, self.base / "result", overwrite=overwrite)

    def test_paper_dimensions_and_original_asset(self):
        report = self.render()
        with Image.open(self.base / "result.png") as im:
            self.assertEqual(im.size, (2126, 1299))
        page = PdfReader(self.base / "result.pdf").pages[0]
        self.assertAlmostEqual(float(page.mediabox.width), 180 / 25.4 * 72, places=3)
        self.assertAlmostEqual(float(page.mediabox.height), 110 / 25.4 * 72, places=3)
        self.assertEqual(hashlib.sha256(self.asset.read_bytes()).hexdigest(), self.original_hash)
        self.assertEqual(report["dimensions"]["width_mm"], 180)
        self.assertAlmostEqual(report["images"][0]["effective_dpi"], 254)
        self.assertEqual(report["images"][0]["placed_bounds_mm"], [10, 25, 60, 30])

    def test_missing_image_rejected_before_output(self):
        self.config["layers"][0]["path"] = "missing.png"
        with self.assertRaises(FileNotFoundError):
            self.render()
        self.assertFalse(list(self.base.glob("result.*")))

    def test_collision_is_rejected_and_explicit_overwrite_is_scoped(self):
        target = self.base / "result.pdf"
        target.write_bytes(b"existing")
        with self.assertRaises(FileExistsError):
            self.render()
        self.assertEqual(target.read_bytes(), b"existing")
        self.assertFalse((self.base / "result.png").exists())
        unrelated = self.base / "keep.txt"
        unrelated.write_text("preserve me")
        self.render(overwrite=True)
        self.assertEqual(unrelated.read_text(), "preserve me")

    def test_pdf_fonts_embedded_and_text_selectable(self):
        self.render()
        page = PdfReader(self.base / "result.pdf").pages[0]
        self.assertIn("Selectable evidence", page.extract_text())
        fonts = page["/Resources"]["/Font"].get_object().values()
        embedded = False
        for ref in fonts:
            font = ref.get_object()
            candidates = font.get("/DescendantFonts", [font])
            for child in candidates:
                descriptor = child.get_object().get("/FontDescriptor")
                if descriptor and "/FontFile2" in descriptor.get_object():
                    embedded = True
        self.assertTrue(embedded)

    def test_svg_text_paths_fonts_and_original_embedded_image(self):
        self.render()
        content = (self.base / "result.svg").read_text()
        root = ET.fromstring(content)
        ns = {"s": "http://www.w3.org/2000/svg"}
        self.assertTrue(root.findall(".//s:text", ns))
        self.assertIn("Selectable evidence", content)
        self.assertIn("data:image/png;base64,", content)
        self.assertIn("data:font/ttf;base64,", content)
        self.assertIn("75,45 83,45 87,30", content)
        self.assertIn('data-layer="math"', content)
        self.assertTrue(root.findall(".//s:path", ns))

    def test_ppt_profile_dimensions_and_editable_shapes(self):
        self.config["profile"] = "ppt"
        report = self.render()
        with Image.open(self.base / "result.png") as im:
            self.assertEqual(im.size, (3840, 2160))
        deck = Presentation(self.base / "result.pptx")
        self.assertAlmostEqual(deck.slide_width / 914400, 13.3333333333, places=4)
        self.assertEqual(deck.slide_height / 914400, 7.5)
        shapes = deck.slides[0].shapes
        self.assertTrue(any(s.has_text_frame and "Selectable evidence" in s.text for s in shapes))
        xml = deck.slides[0]._element.xml
        self.assertIn("cxnSp", xml)
        self.assertIn("tailEnd", xml)
        self.assertTrue(report["limitations"])

    def test_ppt_native_textboxes_have_portable_width_and_line_height(self):
        self.config = {
            "profile": "ppt",
            "layers": [
                {
                    "type": "text",
                    "text": "Field height estimation",
                    "x": 15,
                    "y": 10,
                    "font": "title",
                    "font_size": 30,
                },
                {
                    "type": "text",
                    "text": "Align",
                    "x": 80,
                    "y": 60,
                    "font": "hand",
                    "font_size": 24,
                },
                {
                    "type": "text",
                    "text": "Review",
                    "x": 180,
                    "y": 60,
                    "font": "hand",
                    "font_size": 24,
                },
                {
                    "type": "text",
                    "text": "Fail: realign",
                    "x": 240,
                    "y": 95,
                    "font": "hand",
                    "font_size": 24,
                },
                {
                    "type": "text",
                    "text": "Keep explicit\nwrapped lines",
                    "x": 260,
                    "y": 130,
                    "font": "body",
                    "font_size": 18,
                    "max_width": 60,
                },
            ],
        }
        self.render()
        deck = Presentation(self.base / "result.pptx")
        self.assertEqual(len(deck.slides), 1)
        for shape, layer in zip(deck.slides[0].shapes, self.config["layers"]):
            available = deck.slide_width - shape.left
            self.assertLessEqual(shape.width, available)
            if "max_width" not in layer:
                self.assertAlmostEqual(shape.width / 36000, available / 36000, places=3)
                self.assertFalse(shape.text_frame.word_wrap)
            else:
                self.assertAlmostEqual(shape.width / 36000, layer["max_width"], places=3)
                self.assertTrue(shape.text_frame.word_wrap)
                self.assertEqual(len(shape.text_frame.paragraphs), 2)
            em_mm = layer["font_size"] * 25.4 / 72
            self.assertGreaterEqual(
                shape.height / 36000, len(shape.text_frame.paragraphs) * 1.3 * em_mm
            )
            for paragraph in shape.text_frame.paragraphs:
                self.assertEqual(
                    paragraph.font.name,
                    "Kalam"
                    if layer["font"] == "title"
                    else "Patrick Hand"
                    if layer["font"] == "hand"
                    else "DejaVu Sans",
                )
                self.assertAlmostEqual(paragraph.line_spacing.pt, layer["font_size"] * 1.3)

    def test_ppt_lines_and_rectangles_explicitly_disable_theme_shadows(self):
        self.config["profile"] = "ppt"
        self.render()
        slide = Presentation(self.base / "result.pptx").slides[0]
        connectors = slide._element.xpath(".//p:cxnSp")
        panels = slide._element.xpath(
            './/p:sp[p:spPr/a:prstGeom[@prst="rect"] and not(p:nvSpPr/p:cNvSpPr[@txBox="1"])]'
        )
        self.assertTrue(connectors)
        self.assertTrue(panels)
        for element in connectors + panels:
            self.assertTrue(element.xpath("./p:spPr/a:effectLst"))
            self.assertFalse(element.xpath("./p:spPr/a:effectLst/*"))
            references = element.xpath("./p:style/a:effectRef")
            self.assertTrue(all(ref.get("idx") == "0" for ref in references))

    def test_text_bounds_and_warning(self):
        self.config["layers"].append(
            {"type": "text", "text": "Outside", "font": "body", "x": 178, "y": 106}
        )
        report = self.render()
        self.assertTrue(report["bounds_warnings"])
        self.assertTrue(all(len(item["bounds_mm"]) == 4 for item in report["texts"]))
        saved = json.loads((self.base / "result.render-report.json").read_text())
        self.assertEqual(report, saved)

    def test_math_svg_has_self_contained_vector_glyphs(self):
        self.render()
        root = ET.parse(self.base / "result.svg").getroot()
        ns = {"s": "http://www.w3.org/2000/svg"}
        formula = root.find('.//s:g[@data-layer="math"]', ns)
        self.assertIsNotNone(formula)
        self.assertTrue(
            formula.findall(".//s:path", ns), "math glyphs must be portable vector paths"
        )
        self.assertFalse(
            formula.findall(".//s:text", ns),
            "math fonts must not depend on external CM font installation",
        )

    def test_invalid_geometry_or_formula_produces_no_outputs(self):
        for value in (-1, 0, float("nan")):
            config = dict(self.config, width_mm=value)
            with self.assertRaises(ValueError):
                self.render(config)
            self.assertFalse(list(self.base.glob("result.*")))
        self.config["layers"][3]["text"] = r"$\\definitelyinvalid{x}$"
        with self.assertRaises(ValueError):
            self.render()
        self.assertFalse(list(self.base.glob("result.*")))

    def test_pdf_rebuild_has_fixed_dates_and_identical_bytes(self):
        self.render()
        before = (self.base / "result.pdf").read_bytes()
        self.render(overwrite=True)
        after = (self.base / "result.pdf").read_bytes()
        self.assertEqual(before, after)
        metadata = PdfReader(self.base / "result.pdf").metadata
        self.assertTrue(metadata.creation_date)
        self.assertEqual(metadata.creation_date, metadata.modification_date)

    def test_svg_rebuild_is_identical(self):
        self.render()
        before = (self.base / "result.svg").read_bytes()
        self.render(overwrite=True)
        self.assertEqual(before, (self.base / "result.svg").read_bytes())

    def test_chart_context_uses_assets_without_changing_coordinates(self):
        self.assertTrue(
            (SCRIPTS / "chart_style.py").exists(), "chart style implementation is missing"
        )
        import matplotlib

        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        from chart_style import chart_context, load_style, font_paths

        style = load_style()
        self.assertIn("palette", style)
        self.assertTrue(all(p.exists() for p in font_paths().values()))
        previous = matplotlib.rcParams["axes.prop_cycle"]
        with chart_context():
            fig, ax = plt.subplots()
            (line,) = ax.plot([0, 1, 2], [1.2, 4.7, 3.1])
            self.assertEqual(line.get_xdata().tolist(), [0, 1, 2])
            self.assertEqual(line.get_ydata().tolist(), [1.2, 4.7, 3.1])
            self.assertIsNone(line.get_sketch_params())
            plt.close(fig)
        self.assertEqual(matplotlib.rcParams["axes.prop_cycle"], previous)


if __name__ == "__main__":
    unittest.main()
