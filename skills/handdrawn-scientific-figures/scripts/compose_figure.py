"""Compose ordered scientific figure layers offline using millimetre coordinates.

CLI: compose_figure.py FIGURE.json --output-prefix DEST [--overwrite]
Profiles: paper (180 x 110 mm, 300 dpi) and ppt (16:9, 3840 x 2160).
Text uses x/y visible-top position, font_size in points, font aliases
title/hand/body/bold, optional max_width in mm. Lines/arrows use exact points
and point widths; rough applies only to decorative rect outlines.
Outputs: PNG, embedded-font PDF, self-contained SVG, render-report JSON;
ppt also exports editable text and connectors in a single-slide PPTX.
"""

import argparse
import base64
from datetime import datetime, timezone
from html import escape
from io import BytesIO
import json
import math
import mimetypes
import os
from pathlib import Path
import tempfile
import xml.etree.ElementTree as ET

import matplotlib

matplotlib.use("Agg")
from matplotlib.backends.backend_agg import FigureCanvasAgg
from matplotlib.figure import Figure
from matplotlib.font_manager import FontProperties
from matplotlib import mathtext
from matplotlib.patches import FancyArrowPatch, PathPatch
from matplotlib.path import Path as MplPath
from matplotlib.textpath import TextPath
from PIL import Image

from chart_style import chart_context, font_paths, load_style

PT_MM = 25.4 / 72


def _number(value, label, positive=False):
    value = float(value)
    if not math.isfinite(value) or (positive and value <= 0):
        raise ValueError(f"Invalid {label}: finite{' positive' if positive else ''} value required")
    return value


def _dimensions(config, style):
    profile = config.get("profile", "paper")
    if profile not in ("paper", "ppt"):
        raise ValueError("profile must be paper or ppt")
    defaults = style["profiles"][profile]
    w = _number(config.get("width_mm", defaults["width_mm"]), "width_mm", True)
    h = _number(config.get("height_mm", defaults["height_mm"]), "height_mm", True)
    dpi = _number(config.get("dpi", defaults.get("dpi", 288)), "dpi", True)
    pixels = [round(w / 25.4 * dpi), round(h / 25.4 * dpi)]
    if profile == "ppt" and not any(k in config for k in ("width_mm", "height_mm", "dpi")):
        pixels = defaults["pixels"]
    return profile, w, h, dpi, pixels


def _text_lines(text, prop, max_width):
    lines = []
    for paragraph in text.split("\n"):
        if max_width is None or not paragraph:
            lines.append(paragraph)
            continue
        current = ""
        for word in paragraph.split():
            candidate = (current + " " + word).strip()
            width = TextPath((0, 0), candidate, prop=prop).get_extents().width * PT_MM
            if current and width > max_width:
                lines.append(current)
                current = word
            else:
                current = candidate
        lines.append(current)
    return lines


def _prepare(config, manifest, style):
    """Validate every resource and formula before creating any output."""
    prepared = []
    paths = font_paths(style)
    for p in paths.values():
        if not p.is_file():
            raise FileNotFoundError(f"Missing font resource: {p.name}")
    layers = config.get("layers", [])
    if not isinstance(layers, list):
        raise ValueError("layers must be an ordered array")
    for index, source in enumerate(layers):
        layer = dict(source)
        kind = layer.get("type")
        if kind not in ("image", "text", "math", "arrow", "line", "rect"):
            raise ValueError(f"Unsupported layer {index} type: {kind}")
        layer["index"] = index
        for key in ("x", "y"):
            if kind not in ("arrow", "line"):
                layer[key] = _number(layer[key], key)
        if kind in ("image", "rect"):
            for key in ("width", "height"):
                layer[key] = _number(layer[key], key, True)
        if kind == "image":
            path = manifest.parent / layer["path"]
            if not path.is_file():
                raise FileNotFoundError(f"Missing image resource: {layer['path']}")
            with Image.open(path) as im:
                im.load()
                layer["image"] = im.copy()
                layer["pixel_size"] = im.size
            layer["image_path"] = path
            iw, ih = layer["pixel_size"]
            scale = min(layer["width"] / iw, layer["height"] / ih)
            w, h = iw * scale, ih * scale
            layer["bounds"] = [
                layer["x"] + (layer["width"] - w) / 2,
                layer["y"] + (layer["height"] - h) / 2,
                w,
                h,
            ]
        if kind in ("text", "math"):
            alias = layer.get("font", "body")
            if alias not in paths:
                raise ValueError(f"Unknown font alias: {alias}")
            layer["font"] = alias
            layer["font_size"] = _number(layer.get("font_size", 11), "font_size", True)
            layer["prop"] = FontProperties(fname=str(paths[alias]), size=layer["font_size"])
            if not isinstance(layer["text"], str):
                raise ValueError("text must be a string")
            text = layer["text"]
            if kind == "math":
                text = text if text.startswith("$") and text.endswith("$") else f"${text}$"
                layer["text"] = text
                svg = BytesIO()
                with matplotlib.rc_context({"svg.fonttype": "path"}):
                    mathtext.math_to_image(
                        text,
                        svg,
                        prop=layer["prop"],
                        format="svg",
                        color=layer.get("color", style["palette"]["ink"]),
                    )
                layer["math_svg"] = svg.getvalue().decode()
                lines = [text]
            else:
                lines = _text_lines(text, layer["prop"], layer.get("max_width"))
            layer["text_runs"] = []
            top = layer["y"]
            for line in lines:
                if line:
                    bb = TextPath((0, 0), line, prop=layer["prop"], usetex=False).get_extents()
                    x = layer["x"] - bb.x0 * PT_MM
                    baseline = top + bb.y1 * PT_MM
                    width, height = bb.width * PT_MM, bb.height * PT_MM
                else:
                    x, baseline, width, height = layer["x"], top, 0, layer["font_size"] * PT_MM
                layer["text_runs"].append(
                    {
                        "text": line,
                        "x": x,
                        "baseline": baseline,
                        "bounds": [layer["x"], top, width, height],
                    }
                )
                top += layer["font_size"] * PT_MM * 1.3
            layer["bounds"] = [
                layer["x"],
                layer["y"],
                max(r["bounds"][2] for r in layer["text_runs"]),
                layer["text_runs"][-1]["bounds"][1]
                + layer["text_runs"][-1]["bounds"][3]
                - layer["y"],
            ]
        if kind in ("line", "arrow"):
            points = layer.get("points", [])
            if len(points) < 2 or any(len(p) != 2 for p in points):
                raise ValueError("line and arrow require at least two [x,y] points")
            layer["points"] = [[_number(x, "point x"), _number(y, "point y")] for x, y in points]
            if kind == "arrow" and points[-1] == points[-2]:
                raise ValueError("Arrow final segment must have nonzero length")
        if kind in ("line", "arrow", "rect"):
            layer["line_width"] = _number(
                layer.get("line_width", layer.get("width", 1))
                if kind != "rect"
                else layer.get("line_width", 0.9),
                "line width",
                True,
            )
        prepared.append(layer)
    return prepared


def _rect_points(layer, irregularity):
    x, y, w, h = (layer[k] for k in ("x", "y", "width", "height"))
    corners = [(x, y), (x + w, y), (x + w, y + h), (x, y + h), (x, y)]
    if not layer.get("rough"):
        return corners
    result = []
    for i, (a, b) in enumerate(zip(corners, corners[1:])):
        dx, dy = b[0] - a[0], b[1] - a[1]
        length = math.hypot(dx, dy)
        for j in range(13):
            t = j / 12
            offset = irregularity * math.sin(2 * math.pi * t) * math.sin(math.pi * t)
            result.append(
                (a[0] + t * dx - dy / length * offset, a[1] + t * dy + dx / length * offset)
            )
    result.append(corners[-1])
    return result


def _draw_figure(layers, w, h, dpi, style):
    fig = Figure(figsize=(w / 25.4, h / 25.4), dpi=dpi, facecolor="white")
    FigureCanvasAgg(fig)
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, w)
    ax.set_ylim(h, 0)
    ax.set_axis_off()
    for layer in layers:
        kind, color = layer["type"], layer.get("color", style["palette"]["ink"])
        z = layer["index"] + 1
        if kind == "image":
            x, y, iw, ih = layer["bounds"]
            ax.imshow(layer["image"], extent=(x, x + iw, y + ih, y), aspect="auto", zorder=z)
        elif kind in ("text", "math"):
            for run in layer["text_runs"]:
                ax.text(
                    run["x"],
                    run["baseline"],
                    run["text"],
                    fontproperties=layer["prop"],
                    color=color,
                    va="baseline",
                    zorder=z,
                    parse_math=kind == "math",
                    clip_on=False,
                )
        elif kind in ("line", "arrow"):
            points = layer["points"]
            ax.add_patch(
                PathPatch(
                    MplPath(points),
                    facecolor="none",
                    edgecolor=color,
                    linewidth=layer["line_width"],
                    zorder=z,
                )
            )
            if kind == "arrow":
                ax.add_patch(
                    FancyArrowPatch(
                        points[-2],
                        points[-1],
                        arrowstyle="-|>",
                        shrinkA=0,
                        shrinkB=0,
                        mutation_scale=10,
                        linewidth=layer["line_width"],
                        color=color,
                        zorder=z,
                    )
                )
        else:
            points = _rect_points(layer, style["lines"]["illustrative_irregularity_mm"])
            ax.add_patch(
                PathPatch(
                    MplPath(points),
                    facecolor=layer.get("fill", "none"),
                    edgecolor=layer.get("stroke", color),
                    linewidth=layer["line_width"],
                    linestyle="--" if layer.get("dash") else "-",
                    zorder=z,
                )
            )
    return fig


def _svg(layers, w, h, style):
    fonts = font_paths(style)
    content = [
        f'<svg xmlns="http://www.w3.org/2000/svg" xmlns:xlink="http://www.w3.org/1999/xlink" width="{w:g}mm" height="{h:g}mm" viewBox="0 0 {w:g} {h:g}">',
        "<defs><style>",
    ]
    for alias, path in fonts.items():
        content.append(
            f'@font-face{{font-family:HDSF-{alias};src:url(data:font/ttf;base64,{base64.b64encode(path.read_bytes()).decode()}) format("truetype");}}'
        )
    content += ["</style></defs>", f'<rect width="{w:g}" height="{h:g}" fill="#FFFFFF"/>']
    for layer in layers:
        kind, color = layer["type"], layer.get("color", style["palette"]["ink"])
        content.append(f'<g data-layer="{kind}" id="layer-{layer["index"]}">')
        if kind == "image":
            x, y, iw, ih = layer["bounds"]
            mime = mimetypes.guess_type(layer["image_path"].name)[0] or "image/png"
            raw = layer["image_path"].read_bytes()
            if mime not in ("image/png", "image/jpeg", "image/gif", "image/webp"):
                buffer = BytesIO()
                layer["image"].save(buffer, format="PNG")
                mime, raw = "image/png", buffer.getvalue()
            content.append(
                f'<image x="{x:g}" y="{y:g}" width="{iw:g}" height="{ih:g}" xlink:href="data:{mime};base64,{base64.b64encode(raw).decode()}"/>'
            )
        elif kind == "text":
            for run in layer["text_runs"]:
                content.append(
                    f'<text x="{run["x"]:g}" y="{run["baseline"]:g}" font-family="HDSF-{layer["font"]}" font-size="{layer["font_size"] * PT_MM:g}" fill="{escape(color)}">{escape(run["text"])}</text>'
                )
        elif kind == "math":
            # MathText emits glyph paths and reusable definitions, with no external TeX.
            root = ET.fromstring(layer["math_svg"])
            for metadata in root.findall("{http://www.w3.org/2000/svg}metadata"):
                root.remove(metadata)
            viewbox = root.attrib["viewBox"].split()
            mw, mh = float(viewbox[2]), float(viewbox[3])
            root.set("x", str(layer["x"]))
            root.set("y", str(layer["y"]))
            root.set("width", str(mw * PT_MM))
            root.set("height", str(mh * PT_MM))
            # Namespacing defs avoids colliding glyph IDs across math layers.
            fragment = ET.tostring(root, encoding="unicode")
            import re

            fragment = re.sub(
                r'id="([^"]+)"', lambda m: f'id="math{layer["index"]}-{m[1]}"', fragment
            )
            fragment = re.sub(
                r'(?:href|ns\d+:href)="#([^"]+)"',
                lambda m: f'xlink:href="#math{layer["index"]}-{m[1]}"',
                fragment,
            )
            content.append(fragment)
        elif kind in ("arrow", "line"):
            marker = ""
            if kind == "arrow":
                mid = f"arrowhead-{layer['index']}"
                content.append(
                    f'<defs><marker id="{mid}" markerWidth="7" markerHeight="7" refX="7" refY="3.5" orient="auto" markerUnits="strokeWidth"><path d="M0,0 L7,3.5 L0,7 Z" fill="{escape(color)}"/></marker></defs>'
                )
                marker = f' marker-end="url(#{mid})"'
            points = " ".join(f"{x:g},{y:g}" for x, y in layer["points"])
            content.append(
                f'<polyline points="{points}" fill="none" stroke="{escape(color)}" stroke-width="{layer["line_width"] * PT_MM:g}"{marker}/>'
            )
        else:
            points = _rect_points(layer, style["lines"]["illustrative_irregularity_mm"])
            data = "M" + " L".join(f"{x:g},{y:g}" for x, y in points) + " Z"
            dash = ' stroke-dasharray="2,1.2"' if layer.get("dash") else ""
            content.append(
                f'<path d="{data}" fill="{escape(layer.get("fill", "none"))}" stroke="{escape(layer.get("stroke", color))}" stroke-width="{layer["line_width"] * PT_MM:g}"{dash}/>'
            )
        content.append("</g>")
    return "\n".join(content + ["</svg>"])


def _pptx(layers, w, h, style, output):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.enum.text import MSO_AUTO_SIZE
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Mm, Pt

    deck = Presentation()
    deck.slide_width, deck.slide_height = Mm(w), Mm(h)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)

    def disable_theme_effects(shape):
        # python-pptx's default connector/rectangle styles reference a shadow.
        # Override both the local effect list and the theme reference explicitly.
        for old in shape._element.spPr.xpath("./a:effectLst"):
            shape._element.spPr.remove(old)
        shape._element.spPr.append(OxmlElement("a:effectLst"))
        for ref in shape._element.xpath("./p:style/a:effectRef"):
            ref.set("idx", "0")

    for layer in layers:
        kind, color = layer["type"], layer.get("color", style["palette"]["ink"])
        rgb = RGBColor.from_string(color.lstrip("#"))
        if kind == "image":
            x, y, iw, ih = layer["bounds"]
            slide.shapes.add_picture(str(layer["image_path"]), Mm(x), Mm(y), Mm(iw), Mm(ih))
        elif kind == "text":
            x, y, iw, ih = layer["bounds"]
            # Ink bounds are smaller than an editable text line's layout box,
            # especially when a viewer substitutes an uninstalled hand font.
            # Leave the remaining canvas width available unless the manifest
            # explicitly constrains wrapping, and reserve full em line heights.
            available_width = max(w - x, 1)
            box_width = min(layer.get("max_width", available_width), available_width)
            line_count = len(layer["text_runs"])
            line_height = layer["font_size"] * PT_MM * 1.3
            box_height = min(max(ih + 3, line_count * line_height + 3), max(h - y, 1))
            shape = slide.shapes.add_textbox(Mm(x), Mm(y), Mm(box_width), Mm(box_height))
            tf = shape.text_frame
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            tf.word_wrap = "max_width" in layer
            tf.auto_size = MSO_AUTO_SIZE.NONE
            for i, run in enumerate(layer["text_runs"]):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = run["text"]
                p.font.name = style["fonts"][layer["font"]]["family"]
                p.font.size = Pt(layer["font_size"])
                p.font.bold = layer["font"] in ("title", "bold")
                p.font.color.rgb = rgb
                p.space_before = p.space_after = Pt(0)
                p.line_spacing = Pt(layer["font_size"] * 1.3)
        elif kind == "math":
            buffer = BytesIO()
            mathtext.math_to_image(
                layer["text"], buffer, prop=layer["prop"], dpi=400, format="png", color=color
            )
            buffer.seek(0)
            with Image.open(buffer) as im:
                iw, ih = im.size
            buffer.seek(0)
            slide.shapes.add_picture(
                buffer, Mm(layer["x"]), Mm(layer["y"]), Mm(iw / 400 * 25.4), Mm(ih / 400 * 25.4)
            )
        elif kind in ("line", "arrow"):
            for i, (a, b) in enumerate(zip(layer["points"], layer["points"][1:])):
                shape = slide.shapes.add_connector(
                    MSO_CONNECTOR.STRAIGHT, Mm(a[0]), Mm(a[1]), Mm(b[0]), Mm(b[1])
                )
                shape.line.color.rgb = rgb
                shape.line.width = Pt(layer["line_width"])
                disable_theme_effects(shape)
                if kind == "arrow" and i == len(layer["points"]) - 2:
                    tail = OxmlElement("a:tailEnd")
                    tail.set("type", "triangle")
                    shape.line._get_or_add_ln().append(tail)
        else:
            shape = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Mm(layer["x"]),
                Mm(layer["y"]),
                Mm(layer["width"]),
                Mm(layer["height"]),
            )
            if layer.get("fill") and layer["fill"] != "none":
                shape.fill.solid()
                shape.fill.fore_color.rgb = RGBColor.from_string(layer["fill"].lstrip("#"))
            else:
                shape.fill.background()
            shape.line.color.rgb = RGBColor.from_string(layer.get("stroke", color).lstrip("#"))
            shape.line.width = Pt(layer["line_width"])
            disable_theme_effects(shape)
            if layer.get("dash"):
                shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    stamp = datetime.fromtimestamp(
        int(os.environ.get("SOURCE_DATE_EPOCH", "1791072000")), timezone.utc
    )
    deck.core_properties.created = deck.core_properties.modified = stamp
    deck.core_properties.author = "handdrawn-scientific-figures"
    deck.save(output)


def compose_figure(manifest_path, output_prefix, overwrite=False):
    """Render only the named outputs; fail before output on resource/collision errors."""
    manifest = Path(manifest_path).resolve()
    config = json.loads(manifest.read_text(encoding="utf-8"))
    style = load_style()
    profile, w, h, dpi, pixels = _dimensions(config, style)
    with chart_context(profile):
        layers = _prepare(config, manifest, style)
        prefix = Path(output_prefix)
        suffixes = [".png", ".pdf", ".svg", ".render-report.json"] + (
            [".pptx"] if profile == "ppt" else []
        )
        outputs = {suffix: Path(str(prefix) + suffix) for suffix in suffixes}
        resources = {manifest} | {
            layer["image_path"].resolve() for layer in layers if layer["type"] == "image"
        }
        if any(path.resolve() in resources for path in outputs.values()):
            raise ValueError("Output path collides with a source resource")
        for path in outputs.values():
            if path.exists() and not overwrite:
                raise FileExistsError(f"Output exists: {path.name}; use --overwrite explicitly")
        report = {
            "profile": profile,
            "dimensions": {"width_mm": w, "height_mm": h, "dpi": dpi, "pixels": pixels},
            "outputs": [path.name for path in outputs.values()],
            "images": [],
            "texts": [],
            "bounds_warnings": [],
            "missing_resources": [],
            "errors": [],
            "limitations": ["Raster illustrations remain bitmap assets in all exports."],
        }
        if profile == "ppt":
            report["limitations"] += [
                "PPTX MathText formulas are raster images; PDF/SVG formulas remain vector.",
                "PPTX text is editable but fonts are referenced, not embedded; install bundled fonts on the viewing device.",
                "PPTX decorative rough rectangles use editable straight outlines.",
            ]
        for layer in layers:
            kind = layer["type"]
            if kind == "image":
                iw, ih = layer["pixel_size"]
                bounds = layer["bounds"]
                report["images"].append(
                    {
                        "layer": layer["index"],
                        "path": layer["path"],
                        "source_pixels": [iw, ih],
                        "placed_bounds_mm": bounds,
                        "effective_dpi": min(iw / bounds[2], ih / bounds[3]) * 25.4,
                    }
                )
            if kind in ("text", "math"):
                report["texts"].append(
                    {
                        "layer": layer["index"],
                        "type": kind,
                        "text": layer["text"],
                        "font": layer["font"],
                        "font_size_pt": layer["font_size"],
                        "bounds_mm": layer["bounds"],
                    }
                )
            if "bounds" in layer:
                x, y, bw, bh = layer["bounds"]
            elif kind == "rect":
                x, y, bw, bh = (layer[k] for k in ("x", "y", "width", "height"))
            else:
                xs, ys = zip(*layer["points"])
                x, y, bw, bh = min(xs), min(ys), max(xs) - min(xs), max(ys) - min(ys)
            if x < 0 or y < 0 or x + bw > w or y + bh > h:
                report["bounds_warnings"].append(
                    {
                        "layer": layer["index"],
                        "bounds_mm": [x, y, bw, bh],
                        "message": "Layer extends outside canvas",
                    }
                )
        prefix.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix=".figure-stage-", dir=prefix.parent) as staging:
            stage = Path(staging)
            fig = _draw_figure(layers, w, h, dpi, style)
            buffer = BytesIO()
            fig.savefig(buffer, format="png", dpi=dpi, facecolor="white")
            buffer.seek(0)
            with Image.open(buffer) as image:
                rendered = image.convert("RGB")
                if rendered.size != tuple(pixels):
                    rendered = rendered.resize(tuple(pixels), Image.Resampling.LANCZOS)
                rendered.save(
                    stage / "output.png", dpi=(pixels[0] / (w / 25.4), pixels[1] / (h / 25.4))
                )
            stamp = datetime.fromtimestamp(
                int(os.environ.get("SOURCE_DATE_EPOCH", "1791072000")), timezone.utc
            )
            fig.savefig(
                stage / "output.pdf",
                format="pdf",
                metadata={
                    "CreationDate": stamp,
                    "ModDate": stamp,
                    "Creator": "handdrawn-scientific-figures",
                },
            )
            (stage / "output.svg").write_text(_svg(layers, w, h, style), encoding="utf-8")
            if profile == "ppt":
                _pptx(layers, w, h, style, stage / "output.pptx")
            (stage / "output.render-report.json").write_text(
                json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
            )
            for suffix, output in outputs.items():
                os.replace(stage / ("output" + suffix), output)
            fig.clear()
        return report


def main():
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument(
        "manifest", type=Path, help="JSON manifest; image paths resolve relative to it"
    )
    parser.add_argument(
        "--output-prefix", required=True, type=Path, help="Destination prefix without extension"
    )
    parser.add_argument(
        "--overwrite", action="store_true", help="Replace only this prefix's requested outputs"
    )
    args = parser.parse_args()
    try:
        report = compose_figure(args.manifest, args.output_prefix, args.overwrite)
    except (ValueError, KeyError, FileNotFoundError, FileExistsError) as error:
        parser.exit(2, json.dumps({"error": str(error)}, ensure_ascii=False) + "\n")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
